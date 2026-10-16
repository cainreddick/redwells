#!/usr/bin/env python3
"""Checks the filled request lists with openpyxl: values in the right cells, borders
copied onto extra rows, template formula and image kept, data sheet present."""
import sys, zipfile, pathlib
from openpyxl import load_workbook

folder = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "build/test-output")
problems = []

def check(cond, msg):
    if not cond:
        problems.append(msg)

for n, name in ((6, "request_list_6.xlsx"), (40, "request_list_40.xlsx")):
    path = folder / name
    z = zipfile.ZipFile(path)
    check(z.testzip() is None, f"{name}: zip CRC error")
    check("xl/media/image1.png" in z.namelist(), f"{name}: logo lost")
    wb = load_workbook(path)
    want_sheets = ["Jig Part Request List"] + (["DXF list"] if n == 40 else [])
    check(wb.sheetnames == want_sheets, f"{name}: sheets {wb.sheetnames}")
    check(wb["Jig Part Request List"]["A2"].value == "S180-06068-37005-1", f"{name}: A2 = {wb['Jig Part Request List']['A2'].value!r}")
    check("A2:D2" in [str(r) for r in wb["Jig Part Request List"].merged_cells.ranges], f"{name}: A2:D2 merge lost")
    ws = wb["Jig Part Request List"]
    check(ws["B1"].value == "Jig Part Requst List", f"{name}: title lost")
    check(ws["H1"].value == "=COUNTA(A4:A37)", f"{name}: H1 = {ws['H1'].value!r}")
    check(ws["A3"].value == "Part Number", f"{name}: header lost")
    for i in range(n):
        r = 4 + i
        check(ws[f"A{r}"].value == f"S180-06068-470{i + 2:02d}-1", f"{name}: A{r} = {ws[f'A{r}'].value!r}")
        check(ws[f"B{r}"].value == "MILD STEEL", f"{name}: B{r}")
        check(ws[f"C{r}"].value == ("6mm" if i < n // 2 else "3mm"), f"{name}: C{r}")
        check(ws[f"D{r}"].value == i % 3 + 1, f"{name}: D{r} = {ws[f'D{r}'].value!r} (should be a number)")
        check(ws[f"A{r}"].border.left.style == "thin" and ws[f"D{r}"].border.right.style == "thin", f"{name}: row {r} lost its borders")
    check(ws[f"A{4 + n}"].value is None, f"{name}: extra value after the last row")
    pa = ws.print_area if isinstance(ws.print_area, str) else ",".join(ws.print_area)
    want = 37 if n <= 34 else 3 + n
    check(pa.endswith(f"$D${want}"), f"{name}: print area {pa}, expected to end at row {want}")
    if n == 40:
        ds = wb["DXF list"]
        check([c.value for c in ds[1]] == ["ID", "Revision", "Name", "Qty", "Material", "Thickness", "DXF path"], f"{name}: data header")
        check(ds.max_row == n + 1 and ds["C2"].value == "PLATE 0 & CO <X>", f"{name}: data rows ({ds.max_row}, {ds['C2'].value!r})")
    print(("ok   " if not problems else "FAIL ") + name)

for p in problems:
    print("     " + p)
sys.exit(1 if problems else 0)
