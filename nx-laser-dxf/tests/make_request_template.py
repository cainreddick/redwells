#!/usr/bin/env python3
"""Builds tests/fixtures/request_template.xlsx: a stand-in for the user's
"Jig Part Request List" workbook (logo image, merged title, underlined headers
in row 3, bordered empty rows 4-37, "Total Part No.'s =" with a formula in H1),
used to test filling a real-looking template."""
import pathlib
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Border, Font, Side, Alignment
from PIL import Image as PILImage, ImageDraw

out = pathlib.Path(__file__).resolve().parent / "fixtures"
out.mkdir(exist_ok=True)
logo = out / "logo.png"
img = PILImage.new("RGB", (300, 70), "white")
d = ImageDraw.Draw(img)
d.rectangle([5, 5, 80, 65], fill=(40, 40, 40))
d.text((95, 25), "LOGOS", fill=(200, 30, 30))
img.save(logo)

wb = Workbook()
ws = wb.active
ws.title = "Jig Part Request List"
thin = Side(style="thin")
ws.column_dimensions["A"].width = 36
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 15
ws.column_dimensions["D"].width = 11
ws.row_dimensions[1].height = 58
ws.row_dimensions[2].height = 30
ws.merge_cells("B1:D1")
ws["B1"] = "Jig Part Requst List"
ws["B1"].font = Font(size=16, underline="single")
ws["B1"].alignment = Alignment(horizontal="center", vertical="center")
ws.merge_cells("F1:G1")
ws["F1"] = "Total Part No.'s ="
ws["F1"].font = Font(underline="single")
ws["H1"] = "=COUNTA(A4:A37)"
ws["H1"].alignment = Alignment(horizontal="center")
for col, title in zip("ABCD", ["Part Number", "Material", "Thickness", "Quantity"]):
    c = ws[col + "3"]
    c.value = title
    c.font = Font(underline="single")
    c.alignment = Alignment(horizontal="center")
    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for r in range(4, 38):
    for col in "ABCD":
        ws[col + str(r)].border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws.add_image(Image(str(logo)), "A1")
ws.print_area = "A1:D37"
wb.save(out / "request_template.xlsx")
logo.unlink()
print("wrote", out / "request_template.xlsx")
